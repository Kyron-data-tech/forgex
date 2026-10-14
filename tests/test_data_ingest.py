"""
Unit Tests for DataIngestService (CSV, TSV, XLSX, Preview, Mapping, and Safe Export)
"""

import pytest
import openpyxl
from pathlib import Path

from app.config import settings
from app.services.data_ingest import DataIngestService, escape_formula_cell


@pytest.fixture
def inputs_dir():
    d = settings.DEFAULT_DATA_DIR / "inputs"
    d.mkdir(parents=True, exist_ok=True)
    return d


def test_escape_formula_cell():
    # Dangerous formula characters must be prefixed with single quote
    assert escape_formula_cell("=SUM(A1:A10)") == "'=SUM(A1:A10)"
    assert escape_formula_cell("+12345") == "'+12345"
    assert escape_formula_cell("-2+5") == "'-2+5"
    assert escape_formula_cell("@cmd|' /C calc'!A0") == "'@cmd|' /C calc'!A0"
    assert escape_formula_cell("Normal Text") == "Normal Text"
    assert escape_formula_cell("") == ""
    assert escape_formula_cell(None) == ""


def test_csv_utf8_bom_and_preview(inputs_dir):
    csv_file = inputs_dir / "heroes_bom.csv"
    # Write UTF-8 with BOM
    content = "\ufeffName,Date,Category\nAda Lovelace,1843,Scientists\nMarie Curie,1903,Scientists\n"
    csv_file.write_bytes(content.encode("utf-8"))

    preview = DataIngestService.preview(csv_file)
    assert preview["headers"] == ["Name", "Date", "Category"]
    assert preview["total_rows"] == 2
    assert preview["sample_rows"][0]["Name"] == "Ada Lovelace"
    assert preview["sample_rows"][1]["Name"] == "Marie Curie"


def test_csv_semicolon_delimiter_and_blank_rows(inputs_dir):
    csv_file = inputs_dir / "heroes_semicolon.csv"
    content = "Hero Name;Year;Notes\n\nNelson Mandela;1994;Anti-apartheid\n\nMalala Yousafzai;2014;Nobel Prize\n\n"
    csv_file.write_text(content, encoding="utf-8")

    preview = DataIngestService.preview(csv_file)
    assert preview["headers"] == ["Hero Name", "Year", "Notes"]
    assert preview["total_rows"] == 2


def test_tsv_parsing(inputs_dir):
    tsv_file = inputs_dir / "heroes.tsv"
    content = "Subject\tWorkDate\tType\nAlan Turing\t1936\tHero\nGrace Hopper\t1952\tHero\n"
    tsv_file.write_text(content, encoding="utf-8")

    preview = DataIngestService.preview(tsv_file)
    assert preview["headers"] == ["Subject", "WorkDate", "Type"]
    assert preview["total_rows"] == 2
    assert preview["sample_rows"][0]["Subject"] == "Alan Turing"


def test_excel_multisheet_and_formulas(inputs_dir):
    xlsx_file = inputs_dir / "heroes_multisheet.xlsx"
    wb = openpyxl.Workbook()

    # Sheet 1: Scientists
    ws1 = wb.active
    ws1.title = "Scientists"
    ws1.append(["Person", "Birth Year", "Current Year", "Age Calc"])
    ws1.append(["Nikola Tesla", 1856, 1943, "=C2-B2"])
    ws1.append(["Albert Einstein", 1879, 1955, "=C3-B3"])

    # Sheet 2: Leaders
    ws2 = wb.create_sheet(title="Leaders")
    ws2.append(["Leader Name", "Region"])
    ws2.append(["Mahatma Gandhi", "India"])

    wb.save(xlsx_file)

    sheet_names = DataIngestService.get_sheet_names(xlsx_file)
    assert "Scientists" in sheet_names
    assert "Leaders" in sheet_names

    # Preview Sheet 1
    p1 = DataIngestService.preview(xlsx_file, sheet_name="Scientists")
    assert p1["headers"] == ["Person", "Birth Year", "Current Year", "Age Calc"]
    assert p1["total_rows"] == 2

    # Preview Sheet 2
    p2 = DataIngestService.preview(xlsx_file, sheet_name="Leaders")
    assert p2["headers"] == ["Leader Name", "Region"]
    assert p2["total_rows"] == 1
    assert p2["sample_rows"][0]["Leader Name"] == "Mahatma Gandhi"


def test_column_mapping_and_validation(inputs_dir):
    csv_file = inputs_dir / "validation_test.csv"
    content = (
        "Full_Name,Achievement_Year,Hero_Type,Extra_Info\n"
        "Katherine Johnson,1969,Hero,Apollo 11 trajectory\n"
        ",2000,Hero,Missing name row\n"
        "Chien-Shiung Wu,1956,Hero,Parity violation experiment\n"
    )
    csv_file.write_text(content, encoding="utf-8")

    mapping = {
        "subject_name": "Full_Name",
        "work_date": "Achievement_Year",
        "content_type": "Hero_Type",
    }

    result = DataIngestService.parse_and_validate(csv_file, column_mapping=mapping)
    assert result["total_rows"] == 3
    assert result["valid_count"] == 2
    assert result["error_count"] == 1

    # Verify valid rows have extra context
    first_valid = result["valid_rows"][0]
    assert first_valid["subject_name"] == "Katherine Johnson"
    assert first_valid["work_date"] == "1969"
    assert first_valid["extra_context"]["Extra_Info"] == "Apollo 11 trajectory"

    # Verify error row detected
    err = result["errors"][0]
    assert err["row_number"] == 2
    assert "blank or missing" in err["error"]


def test_path_security_and_file_size():
    # Outside allowlisted directory
    with pytest.raises(PermissionError):
        DataIngestService.validate_file_path(Path("C:/Windows/notepad.exe"))

    # Non-existent file
    with pytest.raises(FileNotFoundError):
        DataIngestService.validate_file_path(settings.DEFAULT_DATA_DIR / "inputs" / "missing.csv")


def test_export_safety_escaping(inputs_dir):
    rows = [
        {"name": "Alice", "formula_field": "=1+1", "note": "@SUM(A1:A5)"},
        {"name": "Bob", "formula_field": "+cmd|'calc'!A0", "note": "-100"},
    ]
    csv_out = inputs_dir / "test_safe_export.csv"
    DataIngestService.export_to_csv(rows, csv_out)

    content = csv_out.read_text(encoding="utf-8-sig")
    assert "'=1+1" in content
    assert "'@SUM(A1:A5)" in content
    assert "'+cmd|'calc'!A0" in content
    assert "'-100" in content

    xlsx_out = inputs_dir / "test_safe_export.xlsx"
    DataIngestService.export_to_xlsx(rows, xlsx_out)

    wb = openpyxl.load_workbook(xlsx_out)
    ws = wb.active
    values = [cell.value for row in ws.iter_rows() for cell in row]
    assert "'=1+1" in values
    assert "'@SUM(A1:A5)" in values
